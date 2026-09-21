# -*- coding: utf-8 -*-
"""Ģenerē Procesu reģistra tehnisko specifikāciju (.docx)."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "Procesu_registrs_tehniska_specifikacija.docx"


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Titullapa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PROCESU REĢISTRS\n")
    r.bold = True
    r.font.size = Pt(22)
    t.add_run("\nTehniskā specifikācija\n\n").font.size = Pt(16)
    t.add_run(f"Versija: 1.0\nDatums: {date.today().strftime('%d.%m.%Y')}\n\n")
    t.add_run("MAPES — Vadības sistēmas arhitektūra")

    doc.add_page_break()

    add_heading(doc, "1. Dokumenta mērķis un apjoms", 1)
    add_para(
        doc,
        "Šis dokuments apraksta «Procesu reģistra» informācijas sistēmas tehnisko arhitektūru, "
        "datu modeli, integrācijas un izvietošanas prasības. Dokuments paredzēts izstrādes, "
        "uzturēšanas un pārcelšanas uz iestādes darba vidi (GitHub, Supabase, statiskā hostēšana) plānošanai.",
    )

    add_heading(doc, "2. Sistēmas kopsavilkums", 1)
    add_para(
        doc,
        "Procesu reģistrs ir tīmekļa lietotne procesu, galaproduktu (GP), procesu jomu, "
        "normatīvo aktu (NA) un izpildītāju uzskaites un pārvaldības atbalstam. "
        "Lietotne ir statiska — HTML, CSS un JavaScript — bez build/soļa un bez atsevišķa backend servera repozitorijā.",
    )
    add_bullets(
        doc,
        [
            "Lietotāja saskarne: viens galvenais fails index.html + modulāri .js faili.",
            "Datu glabāšana: Supabase (PostgreSQL + Storage).",
            "Hostēšana: statisks HTTPS (GitHub Pages, iestādes web serveris u.c.).",
            "Neatbalstīts: atvēršana ar file:// (dubultklikšķis uz HTML).",
        ],
    )

    add_heading(doc, "3. Tehnoloģiju stack", 1)
    add_table(
        doc,
        ["Komponents", "Tehnoloģija", "Piezīmes"],
        [
            ["Frontend", "HTML5, CSS3, Vanilla JavaScript", "Nav React/Vue/Angular"],
            ["Build", "Nav", "Nav package.json, nav bundlera"],
            ["Datu bāze", "Supabase / PostgreSQL", "REST API caur supabase-js"],
            ["Failu glabāšana", "Supabase Storage", "Bucket: pieteikumu-vesture"],
            ["Realtime", "Supabase Realtime + polling", "app:db-sync notikums"],
            ["CDN", "@supabase/supabase-js v2", "jsDelivr / unpkg fallback"],
            ["CI", "GitHub Actions", "Supabase keepalive (nedēļas ping)"],
        ],
    )

    add_heading(doc, "4. Repozitorija un mapju struktūra", 1)
    add_para(doc, "Git repozitorija sakne: vadības sistēmas arhitektura/. Galvenā lietotne atrodas mapē procesu-registrs/.")
    add_bullets(
        doc,
        [
            "procesu-registrs/index.html — lietotnes ieejas punkts",
            "procesu-registrs/DB.js — datu slānis (Supabase)",
            "procesu-registrs/migrations/ — SQL migrācijas",
            "procesu-registrs/docs/ — deploy un specifikācijas dokumentācija",
            "procesu-registrs/config.example.js — Supabase konfigurācijas paraugs",
            "Repozitorija saknē: MAPES kopsavilkuma lapa (index.html), CSV arhīvi, diagrammas",
        ],
    )

    add_heading(doc, "5. Programmatūras moduļi", 1)
    add_table(
        doc,
        ["Modulis", "Atbildība"],
        [
            ["DB.js", "Supabase klients, CRUD, kolonnu mapēšana, Storage, sinhronizācija"],
            ["Procesu registrs.js", "Procesu tabulas datu apvienošana un attēlošana"],
            ["Procesa kartina.js", "Procesa kartiņa, pielikumi"],
            ["GP kataogs.js", "GP kataloga skats, jutīgums"],
            ["Pievienot jaunu procesu.js", "Jauna procesa izveide"],
            ["Joma.js / Joma kartina.js", "Procesu jomas, jomu kartiņa"],
            ["Norm_akti.js", "Normatīvie akti (NA), saistīšana ar procesiem/GP/jomām"],
            ["Izpilditaji.js", "Izpildītāju skats (agregēta tabula)"],
            ["Filtrs.js", "Kolonnu filtri, globālā meklēšana"],
            ["Statistika.js", "Statistikas sadaļas"],
            ["Izmainu_pieteikums.js", "Izmaiņu pieteikums, pielikumi, e-pasts"],
            ["Skaidrojumi.js", "Palīdzības ikonas, BUJ (localStorage)"],
            ["Abreviaturas.js", "Saīsinājumu skaidrojumi"],
            ["Navigacija.js", "Sānu navigācija"],
            ["Numeracija.js", "Procesu/GP numerācijas noteikumi (P/A/M)"],
        ],
    )
    add_para(doc, "Liela daļa biznesa loģikas (tabulu renderēšana, kartiņu redaktori, lomu pārvaldība) ir iekļauta index.html inline skriptā (~3500 rindas).")

    add_heading(doc, "6. Datu modelis", 1)
    add_heading(doc, "6.1. Hierarhija", 2)
    add_para(doc, "Uzdevums → Process(es) → Galaprodukts(i) → (saistības: joma, NA, izpildītājs)")
    add_heading(doc, "6.2. Galvenās entītijas", 2)
    add_table(
        doc,
        ["Entītija", "Glabāšana", "Galvenie lauki / piezīmes"],
        [
            ["Uzdevums", "procesu_registrs", "Uzdevuma_Nr., Uzdevums"],
            ["Process", "procesu_registrs", "Procesa_numurs, Process, Procesa_grupa"],
            ["GP", "procesu_registrs + JSON", "galaprodukti; GP_kartinas_papildu_JSON"],
            ["Joma", "procesu_jomas + Darbibas_joma", "Jomu reģistrs un kartiņas"],
            ["NA", "normativie_akti", "Saistība ar procesu, GP, jomu"],
            ["Klasifikatori", "norm_akti_klasifikatori", "NA veidi, institūcijas"],
        ],
    )
    add_para(doc, "Procesu grupas: pamatdarbība, atbalsts, pārvaldība (prefiksi P / A / M).")
    add_para(doc, "Single-table režīms (noklusējums): GP dati glabājas procesu reģistra tabulā, ne atsevišķā kataloga tabulā.")

    add_heading(doc, "7. Datu bāzes tabulas", 1)
    add_table(
        doc,
        ["Tabula", "Mērķis"],
        [
            ["procesu_registrs", "Galvenā biznesa tabula — procesi, uzdevumi, GP"],
            ["procesu_jomas", "Procesu jomu kartiņas"],
            ["normativie_akti", "Normatīvie akti"],
            ["norm_akti_klasifikatori", "NA klasifikatoru vērtības"],
        ],
    )
    add_para(doc, "Mantots: Procesu_galaproduktu_veidu_katalogs — aizstāts ar single-table modeli (2026-05-05 migrācija).")

    add_heading(doc, "8. SQL migrācijas (secība)", 1)
    add_table(
        doc,
        ["Nr.", "Fails", "Mērķis"],
        [
            ["1", "2026-04-30_joma_un_saglabasana.sql", "Jomas lauki"],
            ["2", "2026-05-05_single-table-procesu-registrs.sql", "Vienas tabulas modelis"],
            ["3", "2026-06-18_procesu_jomas.sql", "Tabula procesu_jomas"],
            ["4", "2026-06-19_procesu_jomas_lauki.sql", "Papildu jomu lauki"],
            ["5", "2026-07-03_gp-nr-to-text.sql", "GP numurs → teksts"],
            ["6", "2026-07-06_gp-kartinas-meta-json.sql", "GP kartiņas meta JSON"],
            ["7", "2026-07-31_normativie_akti.sql", "Normatīvie akti"],
        ],
    )
    add_para(doc, "Migrācijas jāpalaida Supabase SQL Editor secīgi. Detalizēti: migrations/README.md.")

    add_heading(doc, "9. Supabase integrācija", 1)
    add_bullets(
        doc,
        [
            "Savienojums: anon (public) atslēga frontendā — standarta Supabase statiskām lietotnēm.",
            "Konfigurācija: config.js (iestāde) vai noklusējums DB.js (dev vide).",
            "Storage: bucket pieteikumu-vesture — izmaiņu pieteikumu pielikumi, vēsture, kartiņu pielikumi.",
            "Realtime: postgres_changes uz galvenajām tabulām; rezerves polling ik 7 s.",
            "RLS: migrācijās permissive politikas anon/authenticated; produkcijā jāpārskata drošības prasības.",
            "Keepalive: klienta ping + GitHub Actions workflow (free tier neaktivitātes novēršanai).",
        ],
    )

    add_heading(doc, "10. Lietotāju lomas un tiesības", 1)
    add_table(
        doc,
        ["Loma (kods)", "Nosaukums", "Tiesības"],
        [
            ["viewer", "skatītājs", "Tikai skatīšana"],
            ["admin_view", "administrators (skatīt)", "Skatīšana (redaktors bloķēts)"],
            ["admin_edit", "administrators (labot)", "Izveide, labošana, dzēšana"],
        ],
    )
    add_para(
        doc,
        "Lomas glabājas pārlūka localStorage (roleMap). Nav Supabase Auth pieslēguma — "
        "tiesības ir klienta pusē; DB RLS ir galvenais servera līmeņa aizsardzības mehānisms.",
    )
    add_para(doc, "Skata līmeņi (levelSelect): 1 — Pamatskats, 2 — Paplašinātais skats (kolonnu redzamība).")

    add_heading(doc, "11. Galvenās funkcionalitātes", 1)
    add_bullets(
        doc,
        [
            "Procesu reģistrs — tabula, akordeoni, lapošana, Excel eksports",
            "GP katalogs — GP saraksts un paplašinātais skats",
            "Uzdevumu skats — kopsavilkums pa uzdevumiem",
            "Procesu jomas — reģistrs un jomu kartiņas",
            "Procesu grupas — grupēšana un statistika",
            "Izpildītāji — agregēts skats pa pārvaldēm",
            "Normatīvie akti — CRUD, saistīšana ar procesiem/GP/jomām",
            "Statistika — organizācijas, procesu, jomu griezumā",
            "Izmaiņu pieteikums — forma, pielikumi (Storage), e-pasts administratoram",
            "Filtrēšana — kolonnu filtri, globālā meklēšana",
            "Palīdzība — «i» ikonas, BUJ (admin konfigurējams)",
            "Tēmas — gaišs / tumšs režīms",
        ],
    )

    add_heading(doc, "12. localStorage rezerves kopijas", 1)
    add_para(
        doc,
        "Daži moduļi glabā rezerves datus pārlūkā, ja DB tabula nav pieejama vai migrācija nav palaista:",
    )
    add_table(
        doc,
        ["Atslēga / modulis", "Saturs"],
        [
            ["roleMap", "Lietotāju lomas"],
            ["pv_joma_kartinas_v1", "Jomu kartiņas"],
            ["pv_norm_akti_v1", "Normatīvie akti"],
            ["pv_gp_sensitivity_v1", "GP jutīgums"],
            ["pv_help_*", "Palīdzības saturs"],
            ["change_request_log", "Pieteikumu vēsture (teksts)"],
        ],
    )

    add_heading(doc, "13. Konfigurācija un izvietošana", 1)
    add_bullets(
        doc,
        [
            "Kopēt config.example.js → config.js; aizpildīt PV_SUPABASE_URL un PV_SUPABASE_ANON_KEY.",
            "config.js netiek commitots (.gitignore).",
            "Deploy mape: procesu-registrs/ (visi .js blakus index.html).",
            "Pēc git pull — pārlūkā Ctrl+F5 (JS cache bust ar ?v= parametriem).",
            "Detalizēts plāns: docs/DEPLOY_IESTADE.md.",
        ],
    )

    add_heading(doc, "14. Drošība un riski", 1)
    add_table(
        doc,
        ["Risks", "Mitigācija"],
        [
            ["Service role atslēga repozitorijā", "Tikai anon key frontendā; service_role nekad necommitot"],
            ["Klienta lomas", "Produkcijā pārskatīt RLS; nepaļauties tikai uz UI"],
            ["Dev Supabase produkcijā", "Obligāts config.js iestādes projektam"],
            ["Trūkstoša migrācija", "Sekot migrations/README.md secībai"],
            ["file:// atvēršana", "Obligāts HTTP(S) serveris"],
        ],
    )

    add_heading(doc, "15. Nākamie attīstības virzieni (informatīvi)", 1)
    add_bullets(
        doc,
        [
            "Supabase Auth integrācija (īsta lietotāju autentifikācija)",
            "Service role / Edge Functions sensitīvām operācijām",
            "Personu datu bāzes saite izmaiņu pieteikumā",
            "Mērījumu moduļa pilna implementācija",
            "Build pipeline (opcionāli) un automatizēts tests",
        ],
    )

    add_heading(doc, "16. Saistītie dokumenti", 1)
    add_bullets(
        doc,
        [
            "procesu-registrs/README.md",
            "procesu-registrs/docs/DEPLOY_IESTADE.md",
            "procesu-registrs/migrations/README.md",
            "procesu-registrs/SUPABASE_STORAGE_izmainu_pielikumi.md",
            "procesu-registrs/PUBLICET_LASI.md",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("OK:", OUT.name)


if __name__ == "__main__":
    build()
